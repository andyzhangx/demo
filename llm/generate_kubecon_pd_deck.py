#!/usr/bin/env python3
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "kubecon-na-2026-pd-disaggregation.md"
OUT_PATH = ROOT / "kubecon-na-2026-pd-disaggregation.pptx"
ASSET_DIR = ROOT / "deck-assets"
KAITO_LOGO = ASSET_DIR / "kaito-logo.png"
KAITO_DOCS_ARCH = ASSET_DIR / "kaito-docs-arch.png"
KEDA_KAITO_ARCH = ASSET_DIR / "keda-kaito-scaler-arch.png"

TITLE = "Prefill Here, Decode There"
SUBTITLE = "Kubernetes-Native LLM Inference Disaggregation with KAITO and llm-d"
EVENT = "KubeCon + CloudNativeCon North America 2026"

BG = RGBColor(246, 248, 252)
BG_ALT = RGBColor(240, 244, 248)
NAVY = RGBColor(15, 23, 42)
SLATE = RGBColor(51, 65, 85)
MUTED = RGBColor(100, 116, 139)
MID = RGBColor(203, 213, 225)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(37, 99, 235)
BLUE_SOFT = RGBColor(219, 234, 254)
ORANGE = RGBColor(234, 88, 12)
ORANGE_SOFT = RGBColor(255, 237, 213)
GREEN = RGBColor(22, 163, 74)
GREEN_SOFT = RGBColor(220, 252, 231)
PURPLE = RGBColor(109, 40, 217)
PURPLE_SOFT = RGBColor(237, 233, 254)
RED = RGBColor(220, 38, 38)
RED_SOFT = RGBColor(254, 226, 226)
GRAY_SOFT = RGBColor(241, 245, 249)


@dataclass
class SlideSection:
    number: int
    title: str
    lead: List[str]
    bullets: List[str]
    visual: str
    speaker: List[str]


def deck_no(section: SlideSection) -> int:
    return getattr(section, "deck_no", section.number)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, shape_type, x, y, w, h, fill=None, line=None):
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    if fill is None or fill is False:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.color.rgb = MID
        shape.line.width = Pt(1)
    elif line is False:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.25)
    return shape


def add_textbox(
    slide,
    x,
    y,
    w,
    h,
    text="",
    font_size=20,
    bold=False,
    color=NAVY,
    align=PP_ALIGN.LEFT,
    fill=None,
    line=None,
    radius=False,
    margin=0.10,
    font_name="Aptos",
):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    box = add_shape(slide, shape_type, x, y, w, h, fill=fill, line=line)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return box


def add_connector(slide, x1, y1, x2, y2, color=BLUE, width=2.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    return conn


def add_picture(slide, path: Path, x, y, w=None, h=None):
    kwargs = {}
    if w is not None:
        kwargs['width'] = w
    if h is not None:
        kwargs['height'] = h
    return slide.shapes.add_picture(str(path), x, y, **kwargs)


def add_header(slide, title: str, num: int | None = None, accent=BLUE):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.85), fill=NAVY, line=False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0.82), Inches(13.333), Inches(0.04), fill=accent, line=False)
    if num is not None:
        pill = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.20), Inches(0.70), Inches(0.42), fill=WHITE, line=False)
        tf = pill.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"{num:02d}"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = NAVY
        r.font.name = "Aptos"
    add_textbox(slide, Inches(1.28), Inches(0.12), Inches(11.3), Inches(0.55), title, 24, True, WHITE, line=False, fill=False, font_name="Aptos Display")


def clean_bullet(text: str) -> str:
    return re.sub(r"`([^`]+)`", r"\1", text).replace("**", "").strip()


def parse_sections(md_text: str) -> List[SlideSection]:
    pattern = re.compile(r"^### Slide (\d+) — (.+)$", re.MULTILINE)
    matches = list(pattern.finditer(md_text))
    sections: List[SlideSection] = []
    for i, m in enumerate(matches):
        start = m.end()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(md_text)
        body = md_text[start:end].strip()
        title = m.group(2).strip()
        number = int(m.group(1))
        visual = ""
        if "**Suggested visual:**" in body:
            pre, rest = body.split("**Suggested visual:**", 1)
            visual = rest.split("**Speaker note:**", 1)[0].strip().replace("\n", " ")
        else:
            pre = body
        speaker: List[str] = []
        if "**Speaker note:**" in body:
            sp = body.split("**Speaker note:**", 1)[1].strip()
            speaker = [clean_bullet(line[2:]) for line in sp.splitlines() if line.strip().startswith("- ")]
        lead: List[str] = []
        bullets: List[str] = []
        for raw in pre.splitlines():
            line = raw.strip()
            if not line or line.startswith("**"):
                continue
            if line.startswith("- "):
                bullets.append(line[2:].strip())
            else:
                lead.append(line)
        sections.append(SlideSection(number, title, lead, bullets, visual, speaker))
    return sections


def add_card(slide, x, y, w, h, title, body_lines, fill=WHITE, line=MID, title_color=NAVY, body_color=SLATE):
    add_textbox(slide, x, y, w, h, "", fill=fill, line=line, radius=True)
    add_textbox(slide, x + Inches(0.18), y + Inches(0.10), w - Inches(0.36), Inches(0.42), title, 19, True, title_color, line=False, fill=False)
    tb = slide.shapes.add_textbox(x + Inches(0.20), y + Inches(0.58), w - Inches(0.40), h - Inches(0.72))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, line_text in enumerate(body_lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = clean_bullet(line_text)
        p.font.size = Pt(14)
        p.font.color.rgb = body_color
        p.space_after = Pt(5)
    return tb


def add_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, NAVY)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, Inches(0), Inches(13.333), Inches(0.14), fill=ORANGE, line=False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(8.9), Inches(0.14), Inches(4.45), Inches(7.36), fill=RGBColor(22, 32, 58), line=False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(0.90), Inches(0.98), Inches(0.42), fill=WHITE, line=False)
    add_textbox(slide, Inches(0.90), Inches(0.955), Inches(0.82), Inches(0.22), "KubeCon '26", 12, True, NAVY, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(0.82), Inches(1.45), Inches(7.65), Inches(1.15), TITLE, 32, True, WHITE, line=False, fill=False, font_name="Aptos Display")
    add_textbox(slide, Inches(0.84), Inches(2.62), Inches(7.9), Inches(0.86), SUBTITLE, 20, False, RGBColor(201, 214, 255), line=False, fill=False)
    add_textbox(slide, Inches(0.86), Inches(4.10), Inches(5.6), Inches(1.25),
                "A Kubernetes-native story about KAITO, llm-d,\nprefill/decode disaggregation, and inference-aware autoscaling.",
                17, False, WHITE, line=False, fill=False)
    if KAITO_LOGO.exists():
        add_picture(slide, KAITO_LOGO, Inches(6.95), Inches(5.95), h=Inches(0.55))
    chip_y = Inches(5.72)
    chips = [("KAITO", BLUE_SOFT, BLUE), ("llm-d", ORANGE_SOFT, ORANGE), ("GWIE", GREEN_SOFT, GREEN), ("KEDA", PURPLE_SOFT, PURPLE)]
    cx = Inches(0.86)
    for label, fill, line in chips:
        add_textbox(slide, cx, chip_y, Inches(1.22), Inches(0.46), label, 13, True, NAVY, PP_ALIGN.CENTER, fill=fill, line=line, radius=True)
        cx += Inches(1.38)
    add_textbox(slide, Inches(0.86), Inches(6.48), Inches(4.9), Inches(0.45), EVENT, 13, False, RGBColor(201, 214, 255), line=False, fill=False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(9.35), Inches(1.05), Inches(3.15), Inches(5.20), fill=WHITE, line=False)
    add_textbox(slide, Inches(9.65), Inches(1.38), Inches(2.55), Inches(0.42), "Serving topology", 15, True, BLUE, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(9.72), Inches(2.08), Inches(2.40), Inches(0.62), "prefill", 22, True, NAVY, PP_ALIGN.CENTER, fill=ORANGE_SOFT, line=ORANGE, radius=True)
    add_textbox(slide, Inches(9.72), Inches(3.12), Inches(2.40), Inches(0.62), "KV handoff", 22, True, NAVY, PP_ALIGN.CENTER, fill=PURPLE_SOFT, line=PURPLE, radius=True)
    add_textbox(slide, Inches(9.72), Inches(4.16), Inches(2.40), Inches(0.62), "decode", 22, True, NAVY, PP_ALIGN.CENTER, fill=BLUE_SOFT, line=BLUE, radius=True)
    add_connector(slide, Inches(10.92), Inches(2.70), Inches(10.92), Inches(3.12), PURPLE, 3)
    add_connector(slide, Inches(10.92), Inches(3.74), Inches(10.92), Inches(4.16), PURPLE, 3)
    add_textbox(slide, Inches(9.52), Inches(5.18), Inches(2.82), Inches(0.66), "independent scaling\n+ direct KV path", 13, False, MUTED, PP_ALIGN.CENTER, fill=False, line=False)


def add_core_question_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=BLUE)
    add_textbox(slide, Inches(0.85), Inches(1.35), Inches(11.7), Inches(1.25),
                "Why does one LLM request turn into two very different systems problems?", 28, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=False, radius=True)
    add_card(slide, Inches(1.05), Inches(3.0), Inches(4.8), Inches(2.1), "Prefill", ["compute-heavy", "TTFT-sensitive", "prompt bursts hurt first"], fill=ORANGE_SOFT, line=ORANGE)
    add_card(slide, Inches(7.45), Inches(3.0), Inches(4.8), Inches(2.1), "Decode", ["memory / KV heavy", "throughput-sensitive", "generation dominates later"], fill=BLUE_SOFT, line=BLUE)
    add_connector(slide, Inches(5.95), Inches(4.05), Inches(7.35), Inches(4.05), PURPLE, 3)


def add_two_phase_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=ORANGE)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.15), Inches(11.2), Inches(2.1), fill=WHITE, line=False)
    add_textbox(slide, Inches(1.35), Inches(2.62), Inches(4.0), Inches(0.9), "Prefill", 30, True, NAVY, PP_ALIGN.CENTER, fill=ORANGE_SOFT, line=ORANGE, radius=True)
    add_textbox(slide, Inches(8.0), Inches(2.62), Inches(4.0), Inches(0.9), "Decode", 30, True, NAVY, PP_ALIGN.CENTER, fill=BLUE_SOFT, line=BLUE, radius=True)
    add_connector(slide, Inches(5.35), Inches(3.07), Inches(7.95), Inches(3.07), PURPLE, 4)
    add_textbox(slide, Inches(2.05), Inches(4.75), Inches(9.2), Inches(0.8), "One request. Two phases. Two very different bottlenecks.", 24, True, SLATE, PP_ALIGN.CENTER, fill=False, line=False)


def add_kaito_overview_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=GREEN)
    add_textbox(slide, Inches(0.85), Inches(1.20), Inches(11.8), Inches(0.86),
                "KAITO is the Kubernetes control-plane layer that takes you from model serving to distributed inference.",
                22, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=False, radius=True)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.82), Inches(2.20), Inches(4.10), Inches(3.95), fill=WHITE, line=False)
    if KAITO_LOGO.exists():
        add_picture(slide, KAITO_LOGO, Inches(1.05), Inches(2.45), h=Inches(0.60))
    add_textbox(slide, Inches(1.0), Inches(3.20), Inches(3.65), Inches(2.45),
                "• Start with Workspace for one model workload\n\n• Add InferenceSet for replicas and autoscaling\n\n• Add InferencePool for inference-aware routing\n\n• Use MultiRoleInference for richer distributed topology",
                18, False, SLATE, line=False, fill=False)
    if KAITO_DOCS_ARCH.exists():
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(5.15), Inches(2.20), Inches(7.25), Inches(3.95), fill=WHITE, line=False)
        add_picture(slide, KAITO_DOCS_ARCH, Inches(5.35), Inches(2.42), w=Inches(6.85), h=Inches(3.45))
    add_textbox(slide, Inches(0.95), Inches(6.35), Inches(11.4), Inches(0.34),
                "Docs architecture view: KAITO packages serving, routing, and operations into one Kubernetes-native platform model.",
                13, False, MUTED, PP_ALIGN.CENTER, fill=False, line=False)


def add_object_model_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG_ALT)
    add_header(slide, section.title, deck_no(section), accent=PURPLE)
    add_textbox(slide, Inches(0.88), Inches(1.25), Inches(11.7), Inches(0.62),
                "KAITO grows in layers: serving -> scaling -> routing -> distributed topology", 21, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=False, radius=True)
    cols = [
        (Inches(0.95), Inches(2.15), Inches(2.55), Inches(3.95), "Serving", "Workspace", ["Run a model workload", "Serving or tuning", "The basic workload entrypoint"], BLUE_SOFT, BLUE),
        (Inches(3.70), Inches(2.50), Inches(2.55), Inches(3.60), "Scaling", "InferenceSet", ["Replica management", "Autoscaling boundary", "Scale-out for inference"], GREEN_SOFT, GREEN),
        (Inches(6.45), Inches(2.85), Inches(2.55), Inches(3.25), "Routing", "InferencePool", ["Inference-aware routing", "Gateway-facing backend set", "Connects serving to GWIE"], ORANGE_SOFT, ORANGE),
        (Inches(9.20), Inches(3.20), Inches(2.55), Inches(2.90), "Topology", "MultiRoleInference", ["One logical service", "Role-specific backends", "P/D and richer topology"], PURPLE_SOFT, PURPLE),
    ]
    for x, y, w, h, tag, title, lines, fill, line in cols:
        add_textbox(slide, x, y - Inches(0.42), w, Inches(0.32), tag, 12, True, line, PP_ALIGN.CENTER, fill=False, line=False)
        add_card(slide, x, y, w, h, title, lines, fill=WHITE, line=line)
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x + Inches(0.18), y + Inches(0.18), w - Inches(0.36), Inches(0.16), fill=fill, line=False)
    add_connector(slide, Inches(3.50), Inches(4.25), Inches(3.70), Inches(4.25), BLUE, 3)
    add_connector(slide, Inches(6.25), Inches(4.60), Inches(6.45), Inches(4.60), BLUE, 3)
    add_connector(slide, Inches(9.00), Inches(4.95), Inches(9.20), Inches(4.95), BLUE, 3)
    add_textbox(slide, Inches(1.75), Inches(6.42), Inches(9.9), Inches(0.40),
                "The object model stays consistent even as the inference topology becomes more advanced.", 14, False, MUTED, PP_ALIGN.CENTER, fill=False, line=False)


def add_division_of_labor_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=GREEN)
    add_card(slide, Inches(0.9), Inches(2.0), Inches(3.2), Inches(2.4), "KAITO", ["declarative abstraction", "orchestration", "object model"], fill=BLUE_SOFT, line=BLUE)
    add_card(slide, Inches(5.05), Inches(2.0), Inches(3.2), Inches(2.4), "GWIE", ["inference routing contract", "InferencePool model", "Kubernetes-native routing pattern"], fill=GREEN_SOFT, line=GREEN)
    add_card(slide, Inches(9.2), Inches(2.0), Inches(3.2), Inches(2.4), "llm-d Router", ["EPP", "routing + scheduling", "KV/P-D-aware plugins"], fill=ORANGE_SOFT, line=ORANGE)
    add_connector(slide, Inches(4.1), Inches(3.25), Inches(5.0), Inches(3.25), PURPLE, 3)
    add_connector(slide, Inches(8.25), Inches(3.25), Inches(9.15), Inches(3.25), PURPLE, 3)


def add_architecture_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG_ALT)
    add_header(slide, section.title, deck_no(section), accent=ORANGE)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(1.45), Inches(11.85), Inches(4.95), fill=WHITE, line=False)
    add_textbox(slide, Inches(1.05), Inches(1.72), Inches(1.2), Inches(0.28), "request path", 12, True, BLUE, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(1.05), Inches(5.15), Inches(1.2), Inches(0.28), "KV path", 12, True, PURPLE, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(1.05), Inches(2.25), Inches(1.10), Inches(0.78), "Client", 18, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=MID, radius=True)
    add_textbox(slide, Inches(2.35), Inches(2.25), Inches(1.45), Inches(0.78), "Gateway", 18, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=MID, radius=True)
    add_textbox(slide, Inches(4.05), Inches(2.05), Inches(2.15), Inches(1.18), "InferencePool\n+ llm-d EPP", 18, True, NAVY, PP_ALIGN.CENTER, fill=GREEN_SOFT, line=GREEN, radius=True)
    add_textbox(slide, Inches(6.65), Inches(1.55), Inches(2.45), Inches(1.08), "Prefill\nchild InferenceSet", 18, True, NAVY, PP_ALIGN.CENTER, fill=ORANGE_SOFT, line=ORANGE, radius=True)
    add_textbox(slide, Inches(6.65), Inches(3.55), Inches(2.45), Inches(1.08), "Decode\nchild InferenceSet", 18, True, NAVY, PP_ALIGN.CENTER, fill=BLUE_SOFT, line=BLUE, radius=True)
    add_textbox(slide, Inches(9.55), Inches(2.48), Inches(2.30), Inches(1.00), "Routing sidecar\n+ local vLLM", 17, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=PURPLE, radius=True)
    add_connector(slide, Inches(2.15), Inches(2.64), Inches(2.35), Inches(2.64), BLUE, 3)
    add_connector(slide, Inches(3.80), Inches(2.64), Inches(4.05), Inches(2.64), BLUE, 3)
    add_connector(slide, Inches(6.20), Inches(2.40), Inches(6.65), Inches(2.05), BLUE, 2.5)
    add_connector(slide, Inches(6.20), Inches(2.88), Inches(6.65), Inches(4.05), BLUE, 2.5)
    add_connector(slide, Inches(9.10), Inches(4.05), Inches(9.55), Inches(2.98), BLUE, 2.5)
    add_connector(slide, Inches(9.10), Inches(2.05), Inches(9.55), Inches(2.98), PURPLE, 2.5)
    add_textbox(slide, Inches(6.80), Inches(5.15), Inches(4.75), Inches(0.70), "Gateway routes requests; KV data moves directly between prefill and decode workers.", 15, True, SLATE, PP_ALIGN.CENTER, fill=False, line=False)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.00), Inches(4.95), Inches(4.70), Inches(0.72), fill=GRAY_SOFT, line=False)
    add_textbox(slide, Inches(1.15), Inches(5.10), Inches(4.40), Inches(0.36), "Data-plane split: request path and KV-transfer path are different", 13, True, MUTED, PP_ALIGN.CENTER, fill=False, line=False)


def add_compare_signals_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=RED)
    add_textbox(slide, Inches(0.75), Inches(1.30), Inches(12.0), Inches(0.82),
                "CPU-era autoscaling logic came from web workloads, not GPU-bound inference systems.", 24, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=False, radius=True)
    add_card(slide, Inches(0.88), Inches(2.30), Inches(5.25), Inches(3.30), "Classic web-service HPA", [
        "CPU utilization",
        "memory utilization",
        "request-per-pod as a rough proxy",
        "good enough for stateless request/response services",
    ], fill=WHITE, line=MID)
    add_textbox(slide, Inches(5.55), Inches(3.02), Inches(2.22), Inches(1.48), "same CPU\nvery different\nuser pain", 19, True, WHITE, PP_ALIGN.CENTER, fill=RED, line=False, radius=True)
    add_card(slide, Inches(7.18), Inches(2.30), Inches(5.25), Inches(3.30), "LLM inference autoscaling", [
        "waiting requests",
        "token generation pressure",
        "active KV memory",
        "prompt / generation mix",
        "CPU can look fine while TTFT is already bad",
    ], fill=WHITE, line=BLUE)
    add_textbox(slide, Inches(2.25), Inches(6.02), Inches(8.85), Inches(0.46),
                "The autoscaler needs workload-aware signals, not generic resource proxies.", 15, True, SLATE, PP_ALIGN.CENTER, fill=False, line=False)


def add_keda_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=GREEN)
    add_textbox(slide, Inches(0.82), Inches(1.18), Inches(11.75), Inches(0.80),
                "KEDA fits LLM inference because it scales on external workload signals, not just CPU and memory proxies.",
                21, True, NAVY, PP_ALIGN.CENTER, fill=WHITE, line=False, radius=True)
    if KEDA_KAITO_ARCH.exists():
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(2.10), Inches(7.95), Inches(4.15), fill=WHITE, line=False)
        add_picture(slide, KEDA_KAITO_ARCH, Inches(1.15), Inches(2.35), w=Inches(7.55), h=Inches(3.70))
    add_card(slide, Inches(9.15), Inches(2.20), Inches(3.30), Inches(3.95), "Why it matters", [
        "Metric-based scaling for inference queues",
        "Time-based scaling for predictable traffic",
        "InferenceSet as the scaling boundary",
        "Dedicated KAITO scaler removes extra Prometheus dependency for this flow",
    ], fill=WHITE, line=GREEN)


def add_kaito_keda_flow(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG_ALT)
    add_header(slide, section.title, deck_no(section), accent=GREEN)
    labels = [
        ("InferenceSet", BLUE_SOFT, BLUE),
        ("annotations\nscaledobject.kaito.sh/*", WHITE, MID),
        ("KEDA KAITO\nscaler", GREEN_SOFT, GREEN),
        ("ScaledObject\n+ HPA", ORANGE_SOFT, ORANGE),
        ("replicas", WHITE, MID),
    ]
    x = Inches(0.55)
    for i, (label, fill, line) in enumerate(labels):
        bx = x + Inches(2.45) * i
        add_textbox(slide, bx, Inches(2.7), Inches(2.05), Inches(1.35), label, 17, True, NAVY, PP_ALIGN.CENTER, fill=fill, line=line, radius=True)
        if i < len(labels) - 1:
            add_connector(slide, bx + Inches(2.05), Inches(3.38), bx + Inches(2.45), Inches(3.38), BLUE, 3)
    add_textbox(slide, Inches(1.0), Inches(4.85), Inches(11.2), Inches(0.8),
                "Directly scrapes inference metrics from vLLM pods — no separate Prometheus dependency required for this flow.", 17, True, SLATE, PP_ALIGN.CENTER, fill=False, line=False)


def add_role_scaling_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=PURPLE)
    add_card(slide, Inches(0.9), Inches(1.85), Inches(5.3), Inches(3.95), "Prefill autoscaling", [
        "Scale on prompt pressure",
        "Example signal: vllm:num_requests_waiting",
        "Bias toward faster scale-up",
        "Goal: protect TTFT",
    ], fill=ORANGE_SOFT, line=ORANGE)
    add_card(slide, Inches(7.1), Inches(1.85), Inches(5.3), Inches(3.95), "Decode autoscaling", [
        "Scale on sustained generation pressure",
        "Examples: KV/cache pressure, decode saturation",
        "Bias toward throughput stability",
        "Goal: protect sustained token generation",
    ], fill=BLUE_SOFT, line=BLUE)
    add_textbox(slide, Inches(1.55), Inches(6.15), Inches(10.3), Inches(0.55),
                "Once prefill and decode are separate backends, they should scale like separate backends.", 17, True, SLATE, PP_ALIGN.CENTER, fill=False, line=False)


def add_matrix_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, BG_ALT)
    add_header(slide, section.title, deck_no(section), accent=ORANGE)
    add_textbox(slide, Inches(3.95), Inches(1.10), Inches(5.4), Inches(0.50), "KV transfer cost / fabric quality", 18, True, NAVY, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(0.18), Inches(2.95), Inches(0.98), Inches(1.75), "Prefill / decode difference", 15, True, NAVY, PP_ALIGN.CENTER, fill=False, line=False)
    # quadrant labels
    add_textbox(slide, Inches(1.30), Inches(1.35), Inches(5.05), Inches(0.32), "transfer expensive / weak", 13, True, MUTED, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(6.95), Inches(1.35), Inches(5.05), Inches(0.32), "transfer cheap / strong", 13, True, MUTED, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(0.18), Inches(2.35), Inches(0.95), Inches(0.68), "similar", 13, True, MUTED, PP_ALIGN.CENTER, fill=False, line=False)
    add_textbox(slide, Inches(0.10), Inches(4.85), Inches(1.05), Inches(0.92), "meaningfully\ndifferent", 13, True, MUTED, PP_ALIGN.CENTER, fill=False, line=False)
    cells = [
        (Inches(1.30), Inches(1.85), Inches(5.05), Inches(1.95), RED_SOFT, RED, "Stay aggregated", "Small models, short prompts, low concurrency, simple traffic."),
        (Inches(6.95), Inches(1.85), Inches(5.05), Inches(1.95), WHITE, MID, "Usually stay aggregated", "You can try P/D, but the operational gain is often limited."),
        (Inches(1.30), Inches(4.15), Inches(5.05), Inches(1.95), ORANGE_SOFT, ORANGE, "Maybe later", "There may be value, but a weak transfer path will erase it."),
        (Inches(6.95), Inches(4.15), Inches(5.05), Inches(1.95), GREEN_SOFT, GREEN, "Strong candidate for P/D", "Long prompts, retrieval-heavy traffic, long generations, high concurrency, or different scaling / parallelism needs."),
    ]
    for x, y, w, h, fill, line, head, body in cells:
        add_card(slide, x, y, w, h, head, [body], fill=fill, line=line)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(6.88), Inches(4.08), Inches(5.19), Inches(2.09), fill=False, line=GREEN)
    add_textbox(slide, Inches(8.20), Inches(6.32), Inches(2.60), Inches(0.34), "recommended starting point", 12, True, GREEN, PP_ALIGN.CENTER, fill=False, line=False)


def add_generic_slide(prs: Presentation, section: SlideSection):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header(slide, section.title, deck_no(section), accent=BLUE)
    lead = section.lead or []
    bullets = section.bullets or []
    # main white card
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.70), Inches(1.20), Inches(12.0), Inches(5.8), fill=WHITE, line=False)
    tb = slide.shapes.add_textbox(Inches(1.00), Inches(1.55), Inches(10.8), Inches(4.95))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for para in lead:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = clean_bullet(para)
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(10)
    for b in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = clean_bullet(b)
        p.bullet = True
        p.font.size = Pt(21)
        p.font.color.rgb = SLATE
        p.space_after = Pt(8)
    if section.speaker:
        quote = section.speaker[0]
        add_textbox(slide, Inches(8.75), Inches(6.15), Inches(3.15), Inches(0.52), quote, 11, False, MUTED, PP_ALIGN.RIGHT, fill=False, line=False)


def build_presentation(sections: List[SlideSection]) -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_title_slide(prs)
    custom = {
        2: add_core_question_slide,
        3: add_two_phase_slide,
        11: add_kaito_overview_slide,
        12: add_object_model_slide,
        18: add_division_of_labor_slide,
        19: add_architecture_slide,
        26: add_compare_signals_slide,
        27: add_keda_slide,
        28: add_kaito_keda_flow,
        29: add_role_scaling_slide,
        34: add_matrix_slide,
    }
    filtered = [sec for sec in sections if sec.number != 1]
    for idx, sec in enumerate(filtered, 1):
        sec.deck_no = idx
        fn = custom.get(sec.number, add_generic_slide)
        fn(prs, sec)
    return prs


def main():
    md_text = MD_PATH.read_text(encoding="utf-8")
    sections = parse_sections(md_text)
    prs = build_presentation(sections)
    prs.save(str(OUT_PATH))
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
