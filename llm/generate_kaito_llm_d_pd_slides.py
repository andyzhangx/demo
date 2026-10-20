#!/usr/bin/env python3
"""Build the KAITO + llm-d P/D deck using editable PowerPoint shapes.

Requires Python 3.10+ and python-pptx 1.0.2+. Run this file from any directory.
No raster diagrams, external assets, or presentation application are required.
The topology YAML is an explanatory excerpt, not a complete deployment manifest.
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE as SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "kaito-llm-d-pd-slides.pptx"
W, H = 13.333, 7.5
FONT, MONO = "Arial", "Consolas"
NAVY, INK, MUTED = "10243A", "23384D", "5B6E80"
WHITE, BG, LINE = "FFFFFF", "F5F8FC", "DCE5EF"
PURPLE, P_SOFT = "7150C8", "F0ECFA"
BLUE, B_SOFT = "2463D4", "EAF1FF"
AMBER, A_SOFT = "A85C05", "FFF3DD"
TEAL, T_SOFT = "087F82", "E5F5F3"
GREEN, G_SOFT = "24824C", "E9F6EC"
DARK_CARD, LIGHT_TEXT = "19344E", "C8D8E7"
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = "KAITO + llm-d | Prefill/Decode Disaggregation"
prs.core_properties.subject = "Declarative topology, role-aware routing, and direct KV-cache handoff"
prs.core_properties.author = "andyzhangx"
prs.core_properties.keywords = "KAITO, llm-d, GWIE, NIXL, prefill, decode, editable diagrams"
prs.core_properties.comments = "Diagrams are native, editable PowerPoint shapes and connectors."


def rgb(value):
    return RGBColor.from_string(value)


def shape(s, x, y, w, h, fill=WHITE, stroke=None, rounded=True, name=""):
    kind = SHAPE.ROUNDED_RECTANGLE if rounded else SHAPE.RECTANGLE
    obj = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        obj.adjustments[0] = 0.12
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(fill)
    if stroke:
        obj.line.color.rgb = rgb(stroke)
        obj.line.width = Pt(1)
    else:
        obj.line.fill.background()
    obj.name = name or f"Panel {len(s.shapes)}"
    return obj


def text(s, x, y, w, h, value, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.TOP, name=""):
    obj = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    obj.name = name or value.replace("\n", " / ")[:80]
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = valign
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = font
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = rgb(color)
        p.space_before = Pt(0)
        p.space_after = Pt(4 if size >= 18 else 2)
    return obj


def node(s, x, y, w, h, title, sub="", color=BLUE, fill=B_SOFT, size=18):
    obj = shape(s, x, y, w, h, fill, None, name=title)
    tf = obj.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, (value, fs, col, bold) in enumerate([(title, size, color, True)] + ([(sub, 12, MUTED, False)] if sub else [])):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = value
        p.alignment = PP_ALIGN.CENTER
        p.font.name, p.font.size = FONT, Pt(fs)
        p.font.bold, p.font.color.rgb = bold, rgb(col)
        p.space_before, p.space_after = Pt(3 if i else 0), Pt(0)
    return obj


def line(s, x1, y1, x2, y2, color=LINE, width=1.5, arrow=False,
         dashed=False, start=None, end=None, both=False, name="Connector"):
    obj = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                 Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    obj.name = name
    obj.line.color.rgb = rgb(color)
    obj.line.width = Pt(width)
    if dashed:
        obj.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    ln = obj._element.spPr.get_or_add_ln()
    for tag in (["tailEnd"] if arrow else []) + (["headEnd"] if both else []):
        el = OxmlElement(f"a:{tag}")
        el.set("type", "triangle")
        el.set("w", "sm")
        el.set("len", "sm")
        ln.append(el)
    if start:
        obj.begin_connect(*start)
    if end:
        obj.end_connect(*end)
    return obj


def connect(s, a, b, color=BLUE, dashed=False, a_site=3, b_site=1):
    return line(s, 0, 0, 1, 1, color, 1.8, True, dashed,
                (a, a_site), (b, b_site), name=f"{a.name} → {b.name}")


def path(s, points, color=BLUE, dashed=False, width=1.8, name="Flow"):
    for i in range(len(points) - 1):
        line(s, *points[i], *points[i + 1], color, width,
             arrow=i == len(points) - 2, dashed=dashed, name=name)


def pill(s, x, y, w, label, color=BLUE, fill=B_SOFT):
    return node(s, x, y, w, 0.32, label, color=color, fill=fill, size=10)


def dot(s, x, y, label, color=BLUE):
    obj = s.shapes.add_shape(SHAPE.OVAL, Inches(x), Inches(y), Inches(0.29), Inches(0.29))
    obj.name = f"Step {label}"
    obj.fill.solid()
    obj.fill.fore_color.rgb = rgb(color)
    obj.line.fill.background()
    tf = obj.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text, p.alignment = str(label), PP_ALIGN.CENTER
    p.font.name, p.font.size, p.font.bold = FONT, Pt(10), True
    p.font.color.rgb = rgb(WHITE)
    return obj


def footer(s, n, dark=False):
    col = LIGHT_TEXT if dark else MUTED
    text(s, .6, 7.13, 8, .18, "KAITO + llm-d  /  PREFILL–DECODE DISAGGREGATION", 9, col)
    text(s, 11.8, 7.1, .9, .23, f"{n:02d} / 12", 10, col, align=PP_ALIGN.RIGHT)


def slide(n, title, subtitle, section="ARCHITECTURE"):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(WHITE)
    shape(s, .6, .36, .34, .045, TEAL, rounded=False)
    text(s, 1.05, .29, 10.8, .2, f"{section}  /  KAITO + llm-d", 10, MUTED, True)
    text(s, .6, .72, 12.1, .61, title, 30, NAVY, True)
    text(s, .62, 1.39, 12.0, .34, subtitle, 15, MUTED)
    footer(s, n)
    return s


def takeaway(s, value, color=TEAL):
    shape(s, .6, 6.48, 12.13, .49, BG)
    shape(s, .6, 6.48, .045, .49, color, rounded=False)
    text(s, .81, 6.59, 11.7, .26, value, 14, NAVY, True)


def notes(s, value):
    s.notes_slide.notes_text_frame.text = value


def item(s, x, y, number, title, body, color=BLUE, width=3.3):
    dot(s, x, y + .03, number, color)
    text(s, x + .42, y, width - .42, .37, title, 18, NAVY, True)
    text(s, x + .42, y + .49, width - .42, .98, body, 16, MUTED)


def title_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    shape(s, 8.65, 0, 4.683, 7.5, DARK_CARD, rounded=False)
    pill(s, .7, .57, 1.55, "KUBECON ’26", TEAL, T_SOFT)
    text(s, .7, 1.37, 7.8, .8, "KAITO + llm-d", 42, WHITE, True)
    text(s, .72, 2.42, 7.6, 1.5, "Prefill/Decode\nDisaggregation", 36, WHITE, True)
    text(s, .74, 4.15, 7.0, .88, "One declarative topology.\nOne serving endpoint. Two specialized roles.", 22, LIGHT_TEXT)
    pill(s, .74, 5.65, 2.1, "DECLARE TOPOLOGY", PURPLE, P_SOFT)
    pill(s, 3.02, 5.65, 1.94, "ROUTE BY ROLE", BLUE, B_SOFT)
    pill(s, 5.13, 5.65, 2.08, "TRANSFER KV", GREEN, G_SOFT)
    text(s, 9.02, 1.0, 3.5, .3, "THE PLATFORM STACK", 12, LIGHT_TEXT, True)
    labels = [("KAITO", "MultiRoleInference", PURPLE, P_SOFT),
              ("GWIE", "InferencePool contract", BLUE, B_SOFT),
              ("llm-d Router", "Endpoint Picker (EPP)", BLUE, B_SOFT),
              ("vLLM + NIXL", "Runtime + direct KV transfer", TEAL, T_SOFT)]
    boxes = []
    for i, (title, sub, col, fill) in enumerate(labels):
        boxes.append(node(s, 9.02, 1.65 + i * 1.09, 3.58, .86, title, sub, col, fill, 20))
    for a, b in zip(boxes, boxes[1:]):
        connect(s, a, b, LIGHT_TEXT, a_site=2, b_site=0)
    footer(s, 1, True)
    notes(s, "Opening: KAITO owns the control plane; Gateway API Inference Extension (GWIE) defines the gateway-facing contract; llm-d Router supplies the Endpoint Picker (EPP); vLLM and NIXL execute inference and transfer KV cache. This is a responsibility stack, not a network packet path.")


def router_slide():
    s = slide(2, "llm-d Router brings role-aware decisions", "The Endpoint Picker selects workers; the serving path executes the decision.", "ROUTING")
    shape(s, .6, 1.98, 12.13, 1.72, BG)
    a = node(s, .9, 2.32, 2.1, 1.03, "Gateway", "GWIE integration", BLUE, B_SOFT)
    b = node(s, 4.1, 2.24, 3.3, 1.18, "llm-d Router EPP", "Endpoint selection + plugin chain", PURPLE, P_SOFT, 20)
    c = node(s, 8.5, 2.24, 3.8, 1.18, "Selected workers", "Decode + optional prefill", TEAL, T_SOFT, 20)
    connect(s, a, b)
    connect(s, b, c, PURPLE)
    text(s, 3.0, 2.02, 1.2, .2, "pick request", 10, MUTED, align=PP_ALIGN.CENTER)
    text(s, 7.45, 2.02, 1.2, .2, "decision", 10, MUTED, align=PP_ALIGN.CENTER)
    item(s, .8, 4.15, "01", "Select by role", "Model-aware selection\nacross prefill and decode.", BLUE, 3.6)
    item(s, 4.95, 4.15, "02", "Compose plugins", "KV-cache-aware logic\nand P/D-aware scheduling.", PURPLE, 3.6)
    item(s, 9.02, 4.15, "03", "Extend scheduling", "A plugin surface for\nfuture scheduling policy.", TEAL, 3.4)
    takeaway(s, "KAITO builds on llm-d’s routing intelligence — it does not invent a second routing engine.")
    notes(s, "EPP means Endpoint Picker. The arrows show the decision flow, not a proxy hop carrying the inference payload. Endpoint selection may choose a decode worker and, when needed, a prefill worker. Advanced scheduling remains an extension surface, not a claim that every policy is implemented.")


def layers_slide():
    s = slide(3, "Three layers. Three clear responsibilities.", "The InferencePool contract connects platform orchestration to routing.", "RESPONSIBILITIES")
    rows = [(PURPLE, P_SOFT, "KAITO", "DECLARE + RECONCILE", "MultiRoleInference → role topology", "Resource generation • ownership • lifecycle"),
            (BLUE, B_SOFT, "GWIE", "DEFINE THE CONTRACT", "InferencePool + gateway integration", "Kubernetes API surface • backend discovery"),
            (TEAL, T_SOFT, "llm-d Router", "SELECT ENDPOINTS", "EPP + P/D-aware scheduling", "Endpoint picking • role-aware routing decisions")]
    for i, (col, fill, title, tag, main, sub) in enumerate(rows):
        y = 2.01 + i * 1.35
        shape(s, .6, y, 12.13, 1.12, fill)
        text(s, .86, y + .18, 2.7, .38, title, 23, col, True)
        text(s, .88, y + .70, 2.7, .18, tag, 10, col, True)
        line(s, 3.82, y + .2, 3.82, y + .92, col, 1)
        text(s, 4.12, y + .18, 8.2, .35, main, 21, NAVY, True)
        text(s, 4.14, y + .68, 8.0, .25, sub, 15, MUTED)
    takeaway(s, "KAITO owns the topology. GWIE defines the contract. llm-d Router makes the selection.")
    notes(s, "Keep APIs separate from implementations. KAITO uses GWIE objects and integration patterns, but the P/D routing implementation is llm-d Router, not GWIE’s reference EPP implementation. InferencePool is an API object, not an inline traffic proxy.")


def architecture_slide():
    s = slide(4, "One endpoint. Separate control, request and KV paths.", "Read the serving path left to right; endpoint selection is an out-of-band decision.")
    shape(s, .6, 1.98, 12.13, 1.08, P_SOFT)
    text(s, .79, 2.08, 2.3, .18, "CONTROL PLANE", 10, PURPLE, True)
    a = node(s, .8, 2.34, 2.65, .52, "MultiRoleInference", color=PURPLE, fill=WHITE, size=15)
    b = node(s, 4.05, 2.34, 2.6, .52, "KAITO controller", color=PURPLE, fill=WHITE, size=15)
    c = node(s, 7.27, 2.34, 5.2, .52, "Child InferenceSets + InferencePool + wiring", color=PURPLE, fill=WHITE, size=14)
    connect(s, a, b, PURPLE, True)
    connect(s, b, c, PURPLE, True)
    # Routing decision is intentionally above, not in, the request path.
    epp = node(s, 2.25, 3.32, 2.45, .65, "llm-d Router EPP", color=PURPLE, fill=P_SOFT, size=15)
    prefill = node(s, 6.1, 3.32, 2.35, .65, "Prefill vLLM", color=AMBER, fill=A_SOFT, size=16)
    text(s, 5.97, 3.06, 2.6, .19, "PREFILL INFERENCESET", 9, AMBER, True, PP_ALIGN.CENTER)
    shape(s, 5.86, 4.36, 6.61, 1.23, BG, LINE)
    text(s, 6.05, 4.45, 5.8, .19, "DECODE INFERENCESET  /  SELECTED DECODE POD", 9, TEAL, True)
    client = node(s, .65, 4.83, 1.17, .65, "Client", color=BLUE, fill=B_SOFT, size=15)
    gateway = node(s, 2.25, 4.83, 2.45, .65, "Gateway", color=BLUE, fill=B_SOFT, size=16)
    sidecar = node(s, 6.1, 4.83, 2.35, .65, "Routing sidecar", color=BLUE, fill=B_SOFT, size=15)
    decode = node(s, 10.0, 4.83, 2.2, .65, "Decode vLLM", color=TEAL, fill=T_SOFT, size=15)
    connect(s, client, gateway)
    connect(s, gateway, sidecar)
    connect(s, sidecar, decode)
    text(s, 4.75, 4.53, 1.25, .23, "request", 10, BLUE, align=PP_ALIGN.CENTER)
    text(s, 8.58, 4.88, 1.26, .2, "local proxy", 10, BLUE, align=PP_ALIGN.CENTER)
    line(s, 3.48, 4.8, 3.48, 3.98, PURPLE, 1.5, True, True, both=True, name="Gateway ↔ EPP endpoint selection")
    text(s, 1.15, 4.11, 2.08, .44, "endpoint selection\nvia InferencePool", 10, PURPLE, align=PP_ALIGN.RIGHT)
    line(s, 7.28, 4.81, 7.28, 3.99, AMBER, 1.8, True, name="Sidecar → remote prefill request")
    text(s, 5.17, 4.03, 1.9, .22, "optional prefill", 10, AMBER, align=PP_ALIGN.RIGHT)
    path(s, [(8.46, 3.65), (11.1, 3.65), (11.1, 4.81)], GREEN, width=2.3, name="Direct KV data transfer: prefill → decode")
    text(s, 8.85, 3.26, 3.5, .25, "Direct KV data • NIXL", 12, GREEN, True)
    for x, col, label, dash in [(.85, PURPLE, "Control / selection", True), (4.1, BLUE, "Request / local proxy", False), (8.2, GREEN, "KV data (separate path)", False)]:
        line(s, x, 5.97, x + .4, 5.97, col, 2, True, dash)
        text(s, x + .53, 5.85, 3.8, .26, label, 12, MUTED)
    takeaway(s, "The client sees one endpoint; KV data never needs to pass through the gateway.")
    notes(s, "The upper band is resource ownership, not traffic. The gateway consults the llm-d Router EPP through the InferencePool integration and forwards to the selected decode sidecar, carrying the selected prefill worker when applicable. The sidecar invokes remote prefill and local decode. The green arrow denotes the direction of KV data (prefill to decode); the decode runtime initiates the read. Response arrows are omitted for clarity and shown on slide 7. InferenceSets are managed resources, not network hops.")


def topology_slide():
    s = slide(5, "One object declares the serving topology", "Users describe roles and replica intent; controllers own the expansion.", "DECLARATIVE API")
    shape(s, .6, 2.02, 5.23, 4.12, NAVY)
    text(s, .85, 2.23, 4.7, .25, "MultiRoleInference  /  topology excerpt", 12, LIGHT_TEXT, True)
    code = "apiVersion: kaito.sh/v1alpha1\nkind: MultiRoleInference\nmetadata:\n  name: qwen-pd\nspec:\n  roles:\n    - type: prefill\n      replicas: 1\n    - type: decode\n      replicas: 2"
    text(s, .86, 2.75, 4.74, 2.97, code, 16, WHITE, font=MONO)
    text(s, .86, 5.79, 4.65, .18, "Illustrative excerpt — model and placement omitted.", 10, LIGHT_TEXT)
    path(s, [(5.96, 4.07), (6.66, 4.07)], PURPLE)
    text(s, 5.91, 3.62, .79, .3, "reconcile", 10, PURPLE, align=PP_ALIGN.CENTER)
    node(s, 6.9, 2.04, 5.58, .69, "One logical service / InferencePool", color=BLUE, fill=B_SOFT, size=18)
    shape(s, 6.9, 3.05, 2.48, 2.55, A_SOFT)
    shape(s, 9.7, 3.05, 2.78, 2.55, T_SOFT)
    text(s, 7.13, 3.28, 2.05, .35, "Prefill role", 20, AMBER, True)
    text(s, 9.94, 3.28, 2.28, .35, "Decode role", 20, TEAL, True)
    node(s, 7.16, 4.01, 1.95, .71, "Worker 1", color=AMBER, fill=WHITE, size=17)
    node(s, 9.95, 3.99, 2.27, .6, "Worker 1", color=TEAL, fill=WHITE, size=17)
    node(s, 9.95, 4.75, 2.27, .6, "Worker 2", color=TEAL, fill=WHITE, size=17)
    text(s, 7.14, 5.04, 2.04, .24, "replicas: 1", 13, AMBER)
    text(s, 6.94, 5.83, 5.48, .27, "Independent role intent. Shared service surface.", 14, MUTED, True)
    takeaway(s, "The CRD is the product surface — topology intent in, role-specific serving out.")
    notes(s, "This is a topology excerpt, not apply-ready YAML. Model configuration, placement and other required fields are omitted for focus. The original example has one prefill replica and two decode replicas. Replica cards are illustrative workers; an InferenceSet remains the controller-managed abstraction for each role.")


def resources_slide():
    s = slide(6, "KAITO generates resources — and keeps them aligned", "A declarative topology becomes owned resources, runtime wiring and a shared routing contract.", "RECONCILIATION")
    root = node(s, 4.63, 1.98, 4.07, .72, "MultiRoleInference", color=PURPLE, fill=P_SOFT, size=22)
    text(s, 5.0, 2.84, 3.32, .24, "KAITO RECONCILES + OWNS", 11, PURPLE, True, PP_ALIGN.CENTER)
    xs = [.63, 4.83, 9.03]
    colors = [(AMBER, A_SOFT), (TEAL, T_SOFT), (BLUE, B_SOFT)]
    titles = ["Prefill InferenceSet", "Decode InferenceSet", "Routing resources"]
    subs = ["Role-specific workers", "Role-specific workers", "InferencePool + llm-d EPP"]
    for x, (col, fill), title, sub in zip(xs, colors, titles, subs):
        path(s, [(6.665, 3.13), (x + 1.83, 3.13), (x + 1.83, 3.48)], PURPLE, True, name=f"KAITO owns {title}")
        node(s, x, 3.5, 3.67, .93, title, sub, col, fill, 19)
    text(s, .88, 4.71, 3.15, .8, "Role labels\nNIXL host / Pod IP", 17, MUTED)
    text(s, 5.08, 4.71, 3.15, .8, "Routing sidecar injection\nNIXL + local vLLM ports", 17, MUTED)
    text(s, 9.28, 4.71, 3.15, .8, "Services + targetPorts\nEndpoint selection config", 17, MUTED)
    shape(s, .65, 5.74, 12.03, .42, BG)
    text(s, .85, 5.83, 11.6, .22, "Shared invariants: ownership references • role consistency • labels • ports • lifecycle", 13, MUTED, align=PP_ALIGN.CENTER)
    takeaway(s, "Not just generated YAML: controller-owned relationships and consistent runtime contracts.")
    notes(s, "The dashed tree is a resource and responsibility map, not an exact claim that every depicted object is a direct Kubernetes child of the MultiRoleInference. Ownership and reconciliation can be delegated through child controllers. Distinguish the resources from the wiring installed in their pods. Per-role autoscaling is an integration surface; do not imply every policy is provided automatically.")


def sequence_slide():
    s = slide(7, "The decode sidecar coordinates the P/D request", "Remote prefill when needed; local decode on the selected worker.", "REQUEST LIFECYCLE")
    xs = [1.2, 3.63, 6.05, 8.5, 11.2]
    headers = [("Client", BLUE, B_SOFT), ("Gateway", BLUE, B_SOFT), ("Decode sidecar", BLUE, B_SOFT), ("Prefill vLLM", AMBER, A_SOFT), ("Decode vLLM", TEAL, T_SOFT)]
    for x, (label, col, fill) in zip(xs, headers):
        node(s, x - .99, 2.03, 1.98, .54, label, color=col, fill=fill, size=14)
        line(s, x, 2.66, x, 6.31, LINE, 1.1, dashed=True, name=f"{label} lifeline")
    # Small labels immediately above arrows; no text sits across a node or another arrow.
    events = [
        (0, 1, 2.94, "1  Inference request", BLUE, False),
        (1, 2, 3.34, "2  Selected decode + prefill header", BLUE, False),
        (2, 3, 3.74, "3  Remote prefill · max_tokens=1", AMBER, False),
        (3, 2, 4.14, "4  Remote KV parameters", AMBER, True),
        (2, 4, 4.58, "5  Local decode · remote_prefill=true + IDs", BLUE, False),
        (3, 4, 5.02, "6  KV data → decode (NIXL read)", GREEN, False),
        (4, 2, 5.46, "7  Inference response", TEAL, True),
        (2, 1, 5.86, "8  Response", TEAL, True),
        (1, 0, 6.23, "9  Final response", TEAL, True),
    ]
    for a, b, y, label, col, dash in events:
        left, right = sorted([xs[a], xs[b]])
        line(s, xs[a], y, xs[b], y, col, 1.8, True, dash, name=label)
        width = right - left - .12
        # Long adjacent-lane labels use two lines to stay readable at 11 pt.
        if label.startswith("2 "):
            label = "2  Route to decode sidecar\n     + selected prefill header"
        elif label.startswith("3 "):
            label = "3  Remote prefill\n     max_tokens=1"
        elif label.startswith("6 "):
            label = "6  KV data → decode\n     NIXL read by decode"
        two = "\n" in label
        text(s, left + .06, y - (.35 if two else .22), width, .34 if two else .2, label, 10.5 if two else 11, col, bold=False, align=PP_ALIGN.CENTER)
    # Runtime activities sit outside the message labels.
    text(s, 11.44, 5.15, 1.55, .21, "Run decode", 10, TEAL, True)
    text(s, 8.69, 3.84, 1.55, .21, "Run prefill", 10, AMBER, True)
    takeaway(s, "The sidecar coordinates the request; KV data moves directly between the two runtimes.")
    notes(s, "Based on llm-d-router docs/disaggregation.md#pd-sequence. Step 2 follows EPP selection; the EPP consultation is omitted from this sequence and shown on slide 4. Step 4 returns transfer metadata, not the KV payload. Step 5 carries remote_prefill=true, prefill request ID and memory block IDs to local decode. Step 6 shows KV DATA direction (prefill → decode); decode initiates the NIXL read. Response arrows return via sidecar and gateway. Source: https://github.com/llm-d/llm-d-router/blob/main/docs/disaggregation.md#pd-sequence")


def sidecar_slide():
    s = slide(8, "Why the sidecar belongs on the decode pod", "A single decode endpoint coordinates remote prefill and returns the final response.", "POD BOUNDARY")
    shape(s, 4.2, 2.08, 8.48, 2.6, BG, LINE, name="Decode pod boundary")
    text(s, 4.47, 2.29, 7.8, .3, "SELECTED DECODE POD", 12, TEAL, True)
    gateway = node(s, .65, 3.1, 2.58, .95, "Gateway", "EPP-selected endpoint", BLUE, B_SOFT)
    sc = node(s, 4.55, 3.08, 2.85, 1.12, "Routing sidecar", "Port 5000", BLUE, B_SOFT, 20)
    dec = node(s, 9.15, 3.08, 2.95, 1.12, "Local decode vLLM", "Port 5001", TEAL, T_SOFT, 19)
    connect(s, gateway, sc)
    connect(s, sc, dec)
    text(s, 7.52, 3.19, 1.48, .45, "local request\n+ KV metadata", 11, BLUE, align=PP_ALIGN.CENTER)
    pre = node(s, 4.55, 5.17, 2.85, .89, "Remote prefill", "vLLM worker", AMBER, A_SOFT)
    connect(s, sc, pre, AMBER, a_site=2, b_site=0)
    text(s, 2.48, 4.73, 3.16, .25, "optional prefill request", 11, AMBER)
    path(s, [(7.42, 5.61), (10.62, 5.61), (10.62, 4.22)], GREEN, width=2.2, name="NIXL KV data → local decode, bypassing sidecar")
    text(s, 8.05, 5.77, 4.0, .24, "Direct KV data • bypasses sidecar", 12, GREEN, True)
    text(s, .74, 2.34, 2.8, .46, "STABLE EXTERNAL\nROUTING CONTRACT", 11, BLUE, True)
    takeaway(s, "Coordination lives in the sidecar. KV transport terminates at the decode runtime.")
    notes(s, "Ports 5000 and 5001 reflect this deployment configuration, not universal requirements for every llm-d deployment. The InferencePool targetPort must reach the sidecar’s listening port. The sidecar coordinates remote prefill, passes transfer metadata to local vLLM, and proxies the final response. KV payload does not transit through the HTTP sidecar. The decode runtime performs the KV read over NIXL.")


def contracts_slide():
    s = slide(9, "Five contracts the platform must keep consistent", "P/D failures often come from wiring drift — before model execution even begins.", "RUNTIME CONTRACTS")
    cols = [(.86, 3.5), (4.66, 3.44), (8.65, 3.5)]
    for (x, w), label in zip(cols, ["CONTRACT", "REQUIRED INVARIANT", "IF IT DRIFTS"]):
        text(s, x, 2.07, w, .24, label, 11, MUTED, True)
    rows = [
        ("01", "Sidecar placement", "Runs on decode pods", "Coordination path breaks", BLUE),
        ("02", "Port split", "Sidecar :5000 → vLLM :5001", "Proxy loops / wrong listener", BLUE),
        ("03", "Role labels", "Prefill and decode are distinct", "Wrong role gets selected", AMBER),
        ("04", "NIXL reachability", "Advertise reachable Pod IPs", "KV transfer cannot complete", GREEN),
        ("05", "InferencePool targetPort", "Targets the sidecar listener", "Gateway bypasses coordination", PURPLE),
    ]
    for i, (num, title, invariant, failure, col) in enumerate(rows):
        y = 2.53 + i * .69
        shape(s, .62, y, 12.09, .59, BG if i % 2 == 0 else WHITE, rounded=False)
        dot(s, .87, y + .15, num, col)
        text(s, 1.3, y + .15, 3.19, .3, title, 16, NAVY, True)
        text(s, 4.66, y + .16, 3.85, .29, invariant, 15, INK)
        text(s, 8.65, y + .16, 3.86, .29, failure, 15, MUTED)
    takeaway(s, "These are serving contracts, not decorative metadata — reconciliation must enforce them.")
    notes(s, "The right column lists illustrative failure modes, not measured incident statistics. Port values correspond to the illustrated deployment. Reachability includes the correct advertised address and network path. Role labels, selectors and targetPorts must agree with runtime behavior.")


def complement_slide():
    s = slide(10, "KAITO operationalizes llm-d — it does not replace it", "The routing engine and the platform controller solve different problems.", "DIVISION OF LABOR")
    for x, col, fill, title, tag in [(.65, BLUE, B_SOFT, "llm-d Router", "DECIDE WHERE TO SERVE"), (7.0, PURPLE, P_SOFT, "KAITO", "DECLARE + KEEP IT RUNNING")]:
        shape(s, x, 2.05, 5.68, 3.95, fill)
        text(s, x + .28, 2.32, 5.0, .46, title, 28, col, True)
        text(s, x + .29, 2.96, 5.0, .23, tag, 10, col, True)
    left = [("EPP implementation", "Gateway-integrated endpoint picking"), ("P/D-aware scheduling", "Role selection and routing decisions"), ("Extensible plugins", "A surface for advanced scheduling")]
    right = [("Declarative CRD", "Topology and resource abstraction"), ("Resource reconciliation", "Generation, ownership and runtime wiring"), ("Platform integration", "InferenceSet + autoscaling integration")]
    for x, rows in [(.94, left), (7.29, right)]:
        for i, (title, sub) in enumerate(rows):
            y = 3.47 + i * .73
            text(s, x, y, 5.05, .28, title, 18, NAVY, True)
            text(s, x, y + .34, 5.03, .24, sub, 14, MUTED)
    node(s, 6.39, 3.76, .54, .6, "+", color=MUTED, fill=WHITE, size=28)
    takeaway(s, "Complementary layers: routing intelligence below, a Kubernetes-native operating model above.")
    notes(s, "Preserve the boundary: llm-d Router supplies endpoint-picking and scheduling behavior; KAITO supplies declarative intent, resource generation, lifecycle and integration. Autoscaling is shown as an integration point, not a claim of a specific enabled algorithm.")


def ownership_slide():
    s = slide(11, "Users own intent. The platform owns the plumbing.", "A clean operational boundary is more valuable than simply reducing YAML.", "OPERATING MODEL")
    shape(s, .65, 2.04, 4.45, 3.98, P_SOFT)
    shape(s, 6.05, 2.04, 6.63, 3.98, BG)
    text(s, .95, 2.32, 3.85, .47, "Users declare", 25, PURPLE, True)
    text(s, 6.35, 2.32, 6.0, .47, "Platform reconciles", 25, NAVY, True)
    left = ["Model choice", "Prefill / decode roles", "Replicas + autoscaling intent", "Instance types + topology"]
    right = ["Child InferenceSets + shared InferencePool", "llm-d EPP resources + decode sidecar", "Labels, services and targetPorts", "NIXL environment and Pod IP wiring", "Ownership references + per-role consistency"]
    for i, label in enumerate(left):
        dot(s, .98, 3.17 + i * .58, "•", PURPLE)
        text(s, 1.43, 3.18 + i * .58, 3.45, .39, label, 17, INK)
    for i, label in enumerate(right):
        dot(s, 6.38, 3.15 + i * .49, "•", TEAL)
        text(s, 6.83, 3.16 + i * .49, 5.52, .37, label, 17, INK)
    path(s, [(5.22, 4.1), (5.93, 4.1)], PURPLE)
    text(s, 5.17, 4.33, .81, .6, "API\nboundary", 10, PURPLE, align=PP_ALIGN.CENTER)
    takeaway(s, "Users reason about serving behavior — not hidden infrastructure dependencies.")
    notes(s, "Declarative intent is separate from runtime implementation. The platform keeps request handling correct by maintaining these contracts; it cannot eliminate all runtime errors. Autoscaling intent is a policy/integration concern and need not imply that the MultiRoleInference excerpt contains every autoscaling field.")


def close_slide():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY)
    text(s, .72, .51, 11.8, .26, "THE OPERATIONAL PAYOFF", 12, LIGHT_TEXT, True)
    text(s, .72, 1.18, 11.95, 1.5, "From a manual routing stack\nto a declarative platform workflow.", 36, WHITE, True)
    cards = [("01", "Specialize execution", "Prefill and decode have\ndifferent resource profiles.", AMBER, A_SOFT),
             ("02", "Route with awareness", "llm-d adds role-aware\nselection and scheduling.", BLUE, B_SOFT),
             ("03", "Operate declaratively", "KAITO reconciles topology,\nresources and runtime wiring.", TEAL, T_SOFT)]
    for i, (num, title, sub, col, fill) in enumerate(cards):
        x = .72 + 4.2 * i
        shape(s, x, 3.25, 3.89, 2.02, DARK_CARD)
        pill(s, x + .22, 3.47, .57, num, col, fill)
        text(s, x + .22, 4.0, 3.45, .36, title, 21, WHITE, True)
        text(s, x + .22, 4.58, 3.45, .62, sub, 15, LIGHT_TEXT)
    text(s, .74, 5.94, 11.8, .47, "One topology. One endpoint. A repeatable P/D workflow.", 24, WHITE, True)
    footer(s, 12, True)
    notes(s, "Close by restating the complementarity: P/D specializes execution, llm-d supplies routing and scheduling intelligence, and KAITO makes the topology a Kubernetes-native workflow. The value proposition is operational repeatability, not an unsupported performance claim.")


def build():
    for build_slide in [title_slide, router_slide, layers_slide, architecture_slide,
                        topology_slide, resources_slide, sequence_slide, sidecar_slide,
                        contracts_slide, complement_slide, ownership_slide, close_slide]:
        build_slide()
    # Prevent unintended objects outside the canvas; all diagrams remain editable.
    for i, s in enumerate(prs.slides, 1):
        for obj in s.shapes:
            assert obj.left >= 0 and obj.top >= 0, (i, obj.name)
            assert obj.left + obj.width <= prs.slide_width + 20, (i, obj.name)
            assert obj.top + obj.height <= prs.slide_height + 20, (i, obj.name)
    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides; native editable shapes)")


if __name__ == "__main__":
    build()
